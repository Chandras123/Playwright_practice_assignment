from playwright.sync_api import Page
from openpyxl import load_workbook


def test_read_excel_data():
    workbook = load_workbook(
        "test_data/products.xlsx",
        data_only=True,
    )

    sheet = workbook["Products"]

    for row in sheet.iter_rows(
        min_row=2,
        values_only=True,
    ):
        test_id, product, expected_result, status = row

        print("Test ID:", test_id)
        print("Product:", product)
        print("Expected:", expected_result)

    workbook.close()
